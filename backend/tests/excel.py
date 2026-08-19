"""
Actualise verification_toutes_entreprises.xlsx à partir de la table `entreprises`
de la base SQLite du projet.

Usage :
    python actualiser_verification_excel.py

Produit (ou écrase) verification_toutes_entreprises.xlsx dans le dossier courant.
Pensée pour un usage commercial : en-têtes figés, filtres automatiques, et
surlignage des lignes ayant une source d'annuaire non vérifiée (contact à
confirmer manuellement) ou dont le Critic n'a pas validé les sources trouvées.
"""

import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

CHEMIN_DB = "data/dasec_prospection.db"
CHEMIN_SORTIE = "verification_toutes_entreprises.xlsx"

FONT_ENTETE = Font(name="Arial", bold=True, color="FFFFFF")
FILL_ENTETE = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
FONT_NORMAL = Font(name="Arial")
FILL_A_VERIFIER = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # jaune clair
FILL_NON_VALIDE = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")  # rouge clair


def charger_entreprises(chemin_db: str) -> pd.DataFrame:
    if not Path(chemin_db).exists():
        raise FileNotFoundError(f"Base introuvable : {chemin_db}")
    conn = sqlite3.connect(chemin_db)
    try:
        df = pd.read_sql_query("SELECT * FROM entreprises ORDER BY id", conn)
    finally:
        conn.close()
    return df


def exporter_excel(df: pd.DataFrame, chemin_sortie: str):
    df.to_excel(chemin_sortie, index=False, sheet_name="Entreprises")

    from openpyxl import load_workbook
    wb = load_workbook(chemin_sortie)
    ws = wb["Entreprises"]

    # En-têtes
    for cellule in ws[1]:
        cellule.font = FONT_ENTETE
        cellule.fill = FILL_ENTETE
        cellule.alignment = Alignment(horizontal="center", vertical="center")

    # Police par défaut sur toutes les données + repérage des colonnes utiles au surlignage
    entetes = [c.value for c in ws[1]]
    idx_annuaire = entetes.index("contact_annuaire_source") + 1 if "contact_annuaire_source" in entetes else None
    idx_critic = entetes.index("critic_valide") + 1 if "critic_valide" in entetes else None
    idx_telephone = entetes.index("telephone") + 1 if "telephone" in entetes else None
    idx_email = entetes.index("email") + 1 if "email" in entetes else None

    for ligne in ws.iter_rows(min_row=2, max_row=ws.max_row):
        a_verifier = False
        non_valide = False

        if idx_annuaire:
            valeur_annuaire = ligne[idx_annuaire - 1].value
            valeur_tel = ligne[idx_telephone - 1].value if idx_telephone else None
            valeur_email = ligne[idx_email - 1].value if idx_email else None
            # source d'annuaire présente mais aucun contact confirmé => à vérifier manuellement
            if valeur_annuaire and not valeur_tel and not valeur_email:
                a_verifier = True

        if idx_critic:
            valeur_critic = ligne[idx_critic - 1].value
            if valeur_critic in (0, "0", False, "False", None, ""):
                non_valide = True

        for cellule in ligne:
            cellule.font = FONT_NORMAL
            if non_valide:
                cellule.fill = FILL_NON_VALIDE
            elif a_verifier:
                cellule.fill = FILL_A_VERIFIER

    # Largeur des colonnes (approximative, basée sur le contenu)
    for i, colonne in enumerate(ws.columns, start=1):
        longueur_max = max((len(str(c.value)) if c.value is not None else 0) for c in colonne)
        ws.column_dimensions[get_column_letter(i)].width = min(max(longueur_max + 2, 10), 60)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(chemin_sortie)


def main():
    df = charger_entreprises(CHEMIN_DB)
    exporter_excel(df, CHEMIN_SORTIE)
    print(f"{len(df)} entreprises exportées vers {CHEMIN_SORTIE}")
    print("Jaune = source annuaire trouvée mais contact non confirmé (à vérifier manuellement)")
    print("Rouge = sources non validées par le Critic (à revalider)")


if __name__ == "__main__":
    main()